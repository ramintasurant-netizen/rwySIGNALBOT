"""Small helpers on top of python-docx: styles, headings, captions, equations, tables, fields."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, Inches, RGBColor

from eq import render as render_eq

FONT = "Times New Roman"
BODY_PT = 12


def set_run_font(run, size=BODY_PT, bold=None, italic=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    return run


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.54)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    return doc


def para(doc, text="", align="justify", size=BODY_PT, bold=False, italic=False, indent_first=True,
         space_after=6, line_spacing=1.5, left_indent=None, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.keep_with_next = keep_with_next
    if indent_first and align == "justify":
        pf.first_line_indent = Cm(1.0)
    if left_indent is not None:
        pf.left_indent = Cm(left_indent)
    if text:
        add_rich(p, text, size=size, bold=bold, italic=italic)
    return p


def add_rich(p, text, size=BODY_PT, bold=False, italic=False):
    """Very small markup: **bold**, *italic*, ^{sup}, _{sub}."""
    import re
    tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*|\^\{.+?\}|_\{.+?\})", text)
    for t in tokens:
        if not t: continue
        if t.startswith("**"):
            set_run_font(p.add_run(t[2:-2]), size, bold=True, italic=italic)
        elif t.startswith("*"):
            set_run_font(p.add_run(t[1:-1]), size, bold=bold, italic=True)
        elif t.startswith("^{"):
            r = set_run_font(p.add_run(t[2:-1]), size, bold=bold, italic=italic); r.font.superscript = True
        elif t.startswith("_{"):
            r = set_run_font(p.add_run(t[2:-1]), size, bold=bold, italic=italic); r.font.subscript = True
        else:
            set_run_font(p.add_run(t), size, bold=bold, italic=italic)
    return p


def heading1(doc, text, page_break=True):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), 14, bold=True)
    _mark_outline(p, 1)
    return p


def heading2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), 12, bold=True)
    _mark_outline(p, 2)
    return p


def heading3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.left_indent = Cm(0.5)
    set_run_font(p.add_run(text), 12, bold=True, italic=False)
    _mark_outline(p, 3)
    return p


def _mark_outline(p, level):
    ppr = p._p.get_or_add_pPr()
    ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), str(level - 1)); ppr.append(ol)


def numbered(doc, items, start=1, size=BODY_PT, left=1.0, hanging=0.6, fmt="{n}."):
    for i, it in enumerate(items, start):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Cm(left); pf.first_line_indent = Cm(-hanging)
        pf.space_after = Pt(3); pf.line_spacing = 1.5
        pf.tab_stops.add_tab_stop(Cm(left))
        add_rich(p, fmt.format(n=i) + "\t" + it, size=size)


def bullets(doc, items, size=BODY_PT, left=1.0, hanging=0.5, sym="•"):
    for it in items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Cm(left); pf.first_line_indent = Cm(-hanging)
        pf.space_after = Pt(3); pf.line_spacing = 1.5
        pf.tab_stops.add_tab_stop(Cm(left))
        add_rich(p, sym + "\t" + it, size=size)


def equation(doc, tex, number=None, size=12, max_width_in=5.6):
    """Centered display equation (rendered image) with optional right-aligned number."""
    tex = tex.replace(r"\frac", r"\dfrac")  # display-size fractions
    path, w, h = render_eq(tex, fontsize=size)
    if w > max_width_in:
        scale = max_width_in / w; w, h = w * scale, h * scale
    if number is None:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6); p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(path, width=Inches(w))
        return p
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [Cm(13.6), Cm(2.3)]
    for i, (c, wd) in enumerate(zip(t.rows[0].cells, widths)):
        c.width = wd; t.columns[i].width = wd; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c0, c1 = t.rows[0].cells
    p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER; p0.paragraph_format.space_after = Pt(0)
    p0.paragraph_format.line_spacing = 1.0
    p0.add_run().add_picture(path, width=Inches(w))
    p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p1.paragraph_format.space_after = Pt(0)
    set_run_font(p1.add_run(f"({number})"), 12)
    _table_no_borders(t)
    spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = 0.6
    return t


def inline_eq_paragraph(doc, parts, align="left", left_indent=1.0, space_after=4):
    """Paragraph mixing text and small inline formula images: parts = [("t", text) | ("m", tex)]."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if align == "left" else WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Cm(left_indent)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.3
    for kind, val in parts:
        if kind == "t":
            add_rich(p, val)
        else:
            path, w, h = render_eq(val, fontsize=12)
            p.add_run().add_picture(path, width=Inches(w))
    return p


def caption(doc, text, before=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4 if before else 10)
    p.paragraph_format.space_before = Pt(6 if before else 2)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = before
    add_rich(p, text, size=11)
    return p


def picture(doc, path, width_cm, space_after=2):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    return p


def _set_cell_borders(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            el = borders.find(qn(f"w:{edge}"))
            if el is None:
                el = OxmlElement(f"w:{edge}"); borders.append(el)
            for k, v in kwargs[edge].items():
                el.set(qn(f"w:{k}"), str(v))


def _table_no_borders(t):
    for row in t.rows:
        for c in row.cells:
            _set_cell_borders(c, **{e: {"val": "nil"} for e in ("top", "left", "bottom", "right")})


def shade(cell, hex_fill="D9E2F3"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def data_table(doc, header, rows, col_widths_cm, size=11, header_fill="D9E2F3", bold_rows=(), align_first_left=False,
               keep_together=True):
    """Bordered table ('Table Grid'), header bold+shaded, cells centered, rows never split across pages."""
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, w in enumerate(col_widths_cm):
        t.columns[i].width = Cm(w)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.width = Cm(col_widths_cm[i])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        add_rich(p, h, size=size, bold=True)
        shade(c, header_fill)
    for r, row in enumerate(rows, 1):
        for i, val in enumerate(row):
            c = t.rows[r].cells[i]; c.width = Cm(col_widths_cm[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (align_first_left and i == 0) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            add_rich(p, str(val), size=size, bold=(r - 1) in bold_rows)
    # repeat header row on page break; never split a row
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for r_i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        if keep_together and r_i < len(t.rows) - 1:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
    return t


def merge_first_col(table, start_row, end_row):
    a = table.cell(start_row, 0); b = table.cell(end_row, 0)
    a.merge(b)


def add_page_number_footer(section):
    section.different_first_page_header_footer = True
    footer = section.footer
    p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; run._r.append(it)
    set_run_font(run, 11)


def spacer(doc, pt=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(pt); p.paragraph_format.line_spacing = 1.0
    return p
